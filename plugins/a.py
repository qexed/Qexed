from docx import Document

# 创建一个新的Word文档
doc = Document()

# 添加标题
doc.add_heading('HCIP-openEulerV1.0-exam-客观题', level=1)

# 添加考试说明
doc.add_heading('考试说明', level=2)
doc.add_paragraph('分数说明: 总计200分，总成绩占比20%。')
doc.add_paragraph('考试时间: 答题时长:20分钟')
doc.add_paragraph('客观题题型: 总计200分，总成绩占比20%。')
doc.add_paragraph('一、判断题(共28题,合计60.0分)')
doc.add_paragraph('二、单选题(共31题,合计60.0分)')
doc.add_paragraph('三、多选题(共30题，合计60.0分)')
doc.add_paragraph('四、填空题(共5题，合计20.0分)')

# 添加判断题部分
doc.add_heading('一、判断题', level=2)

questions_judgment = [
    "openEuler启动时会读取fstab中的配置，进行自动挂载。无论挂载NAS、SAN或者Gluster，都要在挂载参数中添加'_netdev'，否则挂载会由于网络服务未启动而失败，最终导致系统启动失败。\nA.正确\nB.错误",
    "在Nginx作为反向代理时，客户端向代理请求资源时使用的端口，可以和服务器提供服务端口不一致。\nA.正确\nB.错误",
    "某企业使用LNMP架构搭建了一套Web服务集群，客户端发起请求时的目的端口为80，服务器端可以使用非80端口给反向代理提供资源。\nA.正确\nB.错误",
    "(华为)在SAN存储中,每个target可以关联多个LUN,一个target可以被多个initiators连接,并且target中的每个LUN可以映射给一个initiator使用。\nA.正确\nB.错误",
    "某企业计划使用Nginx作为反向代理,节省IP地址资源并保护真实web服务器,在配置过程中,需要在Nginx中安装https所需的模块并添加代理 export https_proxy=需要代理的https访问地址,并使其配置生效即可。\nA.正确\nB.错误",
    "(华为)某企业内部业务系统较多，所有业务系统均基于LAMP架构实现，为了在访问时更好的区分不同业务，运维工程师使用了LVS的虚拟主机功能来配置不同的域名对应不同的业务。\nA.正确\nB.错误",
    "(华为)某企业使用LNMP架构在华为云ECS上搭建了一套Web服务集群，主机规格为2vCPUs，内存4GiB,并使用Keepalived为LVS提供高可用配置。当有两个客户端同时请求Web服务时，默认情况下Nginx服务器会根据负载均衡算法逐个进行响应。\nA.正确\nB.错误",
    "某企业的业务访问性能要求较高，工程师为了能够最大限度的提高服务器利用率和业务访问效率，使用Keepalived将多个Nginx服务器连接在一起，以提高业务的整体性能。\nA.正确\nB.错误",
    "某企业有两台Apache服务器，为了实现高可用，两台服务器需要做主备，工程师可以利用Keepalived，使用独立的网段在两台服务器之间搭建心跳网络，实现Apache服务器集群的主备功能。\nA.正确\nB.错误",
    "(华为)为实现最小化的访问控制，建议GlusterFS每创建一个卷就在iptables上新开放一个端口。\nA.正确\nB.错误",
    "绝大部分Ansibleplaybook的任务可以通过shell脚本来实现，如果任务对应的功能是无状态的，如web服务器搭建，建议使用Ansibleplaybook来实现，如果任务对应的功能是有状态的，如firewalld规则配置，建议使用shell脚本实现。\nA.正确\nB.错误",
    "(华为)某大型企业IT设备分布在两个城市，其中有两台Nginx服务器运行在两地，建议工程师使用LVS实现两台服务器的负载均衡。\nA.正确\nB.错误",
    "(华为)在LAMP架构的应用中，为了保证MySQL数据库中的数据不丢失，建议使用主备或主从架构,并将log-bin的日志文件保存在GlusterFS提供的分布式卷中。\nA.正确\nB.错误",
    "NFS服务器提供共享存储时，如果需要向所有用户开放共享目录/share/data的读写权限，需要在配置文件/etc/exports中添加/share/data*(ro,sync)\nA.正确\nB.错误",
    "Nginx和 Apache一样,均采用模块化的设计,支持丰富的第三方模块,并且支持利用Linux的多线程原理实现客户端请求的并发处理。\nA.正确\nB.错误",
    "(华为)在当今IT系统中，几乎所有的业务都是集群化部署，当其中一个业务主机发生故障时，集群内其他主机会自动接管故障主机的任务，用户侧几乎是无感知的，但实际总体性能会下降。为了保证性能迅速恢复，管理员可通过自动化运维工具，如Ansible、SaltStack等，自动对备用空服务器进行上电和安装操作系统，同时将部署应用的shell脚本发送到该主机上并自动运行，迅速使其升级为业务主机。\nA.正确\nB.错误",
    "(华为) Zabbix Server， Zabbix的Web页面，Nginx三者可以在不同的服务器中部署。\nA.正确\nB.错误",
    "某企业IT集群中的服务器性能相近，使用最少连接算法可以进行较好的负载均衡，在该场景中，LVS、Nginx、HAProxy均可以满足需求。\nA.正确\nB.错误",
    "在企业IT架构中，为了提高业务访问效率，可以部署多台服务器，并使用LVS做负载均衡，其中LVS的调度算法包含静态调度算法和动态调度算法，这两种算法均不需要根据工作模式来调整。\nA.正确\nB.错误",
    "某企业使用四台服务器搭建了Nginx集群，在配置高可用时，Keepalived不仅能利用VRRP协议实现两两双机热备，也可以同时热备所有服务器，来满足该企业的需求。\nA.正确\nB.错误",
    "(华为)Glusterfs的默认卷类型为分布式，在该模式下，数据未进行冗余保护，一旦一个节点损坏,数据就会丢失,因此该模式不建议使用在LNMP架构中。\nA.正确\nB.错误",
    "在LNMP架构中,firewalld需要在每个节点中都开启并配置。在Nginx服务器中,需要放通其对应的业务端口，同时需要放通访问数据库的端口。在数据库服务器中，需要放通对应的业务接口即可。\nA.正确\nB.错误",
    "在LAMP架构的应用中,web服务的静态页面被保存在GlusterFS提供的卷中,为了保证性能,管理员通过GlusterFS原生客户端的方式，将卷挂载到web服务器指定目录下。另外，保证web服务器重启后能够自动挂载该卷,管理员采用了编写相应的 shell脚本,使其开机自动运行,作为最佳实现方式。\nA.正确\nB.错误",
    "某业务使用LVS做负载均衡，其对外虚拟IP为192.168.1.21端口号为80，DIP为172.16.1.21，实际业务主机地址为172.16.1.33-37,端口号为8087。某用户IP为192.168.100.100访问业务,向LVS发起请求。请求被分发到实际业务主机时，业务主机收到的请求报文的目的地址为172.16.1.21，端口号为 8087。\nA.正确\nB.错误",
    "(华为)某企业计划搭建web服务集群，为了实现web服务的高可用，可以通过Nginx提供的七层负载均衡功能来实现。\nA.正确\nB.错误",
    "在一台Apache服务器上运行了多台虚拟主机，每个虚拟主机可以对应多个Web服务，用户可以使用不同的端口号来访问到不同的资源。\nA.正确\nB.错误",
    "(华为)数据库是LAMP中非常重要的部分,保存着全部的业务数据,为了保证数据的可靠性,管理员将SAN存储LUN挂载到不同的节点上,作为Glusterfs的brick,最后再使用brick组成复制卷用来存放数据。\nA.正确\nB.错误",
    "使用命令 `find / -nouser | xargs -i test -s {} && echo $?` 可直找出无主的文件或目录,并确定这些文件是否为空。\nA.正确\nB.错误"
]

for i, question in enumerate(questions_judgment, 1):
    doc.add_paragraph(f"{i}. {question}")

# 添加单选题部分
doc.add_heading('二、单选题', level=2)

questions_single_choice = [
    "某企业在内网新增了多个业务,并通过域名+不同URL的方式来进行访问,其中域名使用的企业固定域名。以下哪个工具可实现该需求?\nA. Keepalived\nB. Apache\nC. LVS\nD. Nginx",
    "以下关于Keepalived和LVS的描述中，错误的是哪一项?\nA. LVS本身只能完成负载均衡的调度\nB. 虚拟IP是Keepalived提供的，而不是LVS本身自带的\nC. Keepalived最初是为LVS设计的\nD. LVS的负载均衡可以避免单点故障",
    "以下哪个选项符合http的工作机制?\nA. 客户向Apache服务器发起请求，Apache服务器通过路径重定向找到对应的资源，将其反馈给用户，同时保存一份到数据库中，进行事务记录\nB. 客户向Apache服务器发起请求，Apache服务器检查firewalld的策略，将符合放行规则的资源反馈给用户\nC. 客户向web服务器发起请求，为web服务器进行负载均衡的LVS会首先对该请求进行响应，并检查客户端所请求的资源是否在本地缓存，如果有，直接反馈给客户，如果没有再向后端的web服务器进行获取，最后反馈给客户。\nD. 客户向web服务器发起请求，web服务器收到请求后，可从其他web服务器上读取和调用所需的资源，然后进行整合，并将其反馈给客户",
    "在某企业中，管理员使用Nginx为三台Apache服务器提供反向代理，被代理的url分别为www.test80.com:80、www.test81.com:81和www.test82.com:82，配置完成后，用户通过www.test.com:80可以访问对应的业务。如果Nginx服务所使用的网卡绑定在public zone，根据最小化的安全策略，管理员应在Nginx上的public zone中放通哪个端口?\nA. 82\nB. 80、81、82和53\nC. 80\nD. 81",
    "某wordpress网站集群挂载GlusterFS的卷作为后端存储，保存网站的静态数据。GlusterFS使用三台主机组成存储池，并使用keepalived的虚拟IP对外提供服务。以下关于此场景下不同类型存储卷的描述中，错误的是哪一项?\nA. 三台主机组成差分卷,其中两台主机损坏, wordpress网站中全部静态数据仍可正常访问\nB. 三台主机组成分布卷，其中一台主机损坏， wordpress网站中部分静态数据将无法正常使用\nC. 三台主机组成分布卷，其中两台主机损坏， wordpress网站中部分静态数据仍可正常访问\nD. 三台主机组成复制卷，其中一台主机损坏， wordpress网站中全部静态数据仍可正常访问",
    "使用Apache搭建网站服务集群，两台服务器的IP地址分别10.0.0.1和10.0.0.2，前端使用Nginx作为负载均衡代理，其地址为192.168.1.1。使用DNS服务进行地址解析，用户访问www.test.com时即可访问该网站。请问在DNS配置中配置哪个解析，可以实现用户需求?\nA. www AAAA 192.168.1.1\nB. www MX 192.168.1.1\nC. www A 192.168.1.1\nD. www PTR 192.168.1.1",
    "Keepalived不适用于以下哪种业务的集群搭建?\nA. web服务集群\nB. Zabbix集群\nC. LVS集群\nD. DNS集群",
    "某企业在进行IT架构设计时，计划使用LVS来实现负载均衡，如果后端服务器数量较多，网段无法统一，以下哪项为合理的设计方案?\nA. 使用 NAT模式对后端服务器进行负载均衡\nB. 将后端服务器的IP地址修改为同一网段，然后使用DR模式进行负载均衡\nC. 使用 DR模式对后端服务器进行负载均衡\nD. 为LVS服务器配置多个网卡和IP地址,根据后端服务的业务量选择合适的模式进行负载均衡",
    "在进行GlusterFS集群搭建时，以下哪个选项中的操作，相比ansible playbook，更适合使用shell脚本完成?\nA. 安装并配置 Keepalived以实现 GlusterFS的高可用\nB. 周期性检查 bricks所在磁盘的容量\nC. 获取 node节点所有磁盘容量\nD. 在DNS服务器中添加 node主机名解析",
    "使用 Nginx作为负载均衡器,将流量转发至后端 Nginx服务集群web01(10.0.0.7)和web02(10.0.0.8)。在实际转发分配中要求在保证所有服务都可以接收到请求的情况下，web01接受更多请求，可以使用以下哪个配置实现?\nA.\n```\nupstream load_pass{\n    ip_hash;\n    server 10.0.0.7:80 weight=5;\n    server 10.0.0.8:80 down;\n}\n```\nB.\n```\nupstream load_pass{\n    server 10.0.0.7:80 weight=5;\n    server 10.0.0.8:80;\n}\n```\nC.\n```\nupstream load_pass{\n    server 10.0.0.7:80;\n    server 10.0.0.8:80;\n}\n```\nD.\n```\nupstream load_pass{\n    server 10.0.0.7:80;\n    server 10.0.0.8:80 backup;\n}\n```",
    "某企业新采购了两台性能较好的服务器，用来替换两台老性能较差易宕机、运行了web服务的旧服务器，工程师为了实现服务的高可用，以下哪一种方案是不可取的?\nA. 两台新服务器借助Keepalived组成web高可用集群\nB. 一台新服务器运行多个虚拟机主机，并组成web集群，另外一台服务器作为负载均衡部署在前端\nC. 两台新服务器使用虚拟机主机组成不同的web集群，使用keepalived为其提供高可用\nD. 将两台新服务器加入到现有业务集群中，逐步接收旧服务器的流量，最后等系统稳定后，再将旧服务器下线",
    "某工程师在使用LVS对web服务进行负载均衡时，为了提高LVS的转发效率，同时避免其不成为整个架构的性能瓶颈，建议使用以下哪种模式?\nA. RR\nB. DR\nC. NAT\nD. TUN",
    "根据下图中CPU的信息，Nginx启动后会默认自动创建多少个worker进程?\nA. 3\nB. 4\nC. 5\nD. 2",
    "某工程师想要使用LVS的NAT工作模式实现负载均衡，VIP为192.168.1.10，DIP为10.0.0.1，其中一台RS的IP为10.0.0.2，则以下哪一条添加RS的配置是正确的?\nA. `ipvsadm -a -t 10.0.0.1:80 -r 10.0.0.2`\nB. `ipvsadm -a -t 192.168.1.10:80 -r 10.0.0.2 -m`\nC. `ipvsadm -a -t 192.168.1.10:80 -r 10.0.0.2`\nD. `ipvsadm -a -t 10.0.0.2:80 -r 10.0.0.1 -m`",
    "某工程师下载了Nginx的源码安装包，在安装完编译依赖环境后，需要使用以下哪一个工具执行编译构建?\nA. Ninja\nB. make\nC. GNU\nD. newlib",
    "黑客通过某种方式入侵了企业内网,现正在尝试对业务系统进行攻击,以下关于攻击方法和现象描述中，错误的是哪一项?\nA. 黑客已知企业正在使用 GlusterFS相关业务功能,因此考虑从udp 111端口进行攻击\nB. 某内部网站使用LVS负载均衡,黑客针对LVS的80端口进行泛洪攻击,达到网站瘫痪目的\nC. 黑客通过端口扫描，发现某一IP启用了3389端口，因此计划采用撞库的方式进行攻击\nD. 在通过ping命令扫描内网设备时,发现所有IP都提示'无法访问目标主机',可能原因是在iptables里配置了REJECT",
    "大型企业需要部署一万台服务器，为了更高效的完成业务的安装部署任务，可以采用自动化运维工具。在该场景中，以下哪一项运维工具能够更好地满足需求?\nA. Ansible\nB. SaltStack\nC. Shell\nD. Puppet",
    "某工程师使用三台性能不同的服务器进行web服务集群的搭建，为了保证各个服务器的性能都能被充分利用，同时还需要保证该集群业务的可用性，以下哪个选项为较合理的设计方案?\nA. 使用一台服务器配置Nginx为另外两台做负载均衡，使用轮询算法为后端应用服务器做负载均衡\nB. 使用一台服务器配置Nginx为另外两台做负载均衡,配置虚拟主机,使用不同端口区分不同服务\nC. 使用一台服务器配置Nginx为另外两台做负载均衡,使用权重算法为后端应用服务器做负载均衡，性能高的权重高\nD. 配置虚拟主机,其中两个虚拟主机位于不同物理服务器上,使用keepalived实现高可用集群,同时这两台虚拟主机使用合理的算法为其他虚拟主机提供负载均衡调度",
    "某工程师配置 Nginx为后端服务器做负载均衡,根据以下配置分析,哪个选项中的服务器响应访问的次数最多?\n```\nupstream xxx{\n    server 192.168.1.11 down;\n    server 192.168.1.12 weight=2;\n    server 192.168.1.13;\n    server 192.168.1.14 backup;\n}\n```\nA. 192.168.1.11\nB. 192.168.1.12\nC. 192.168.1.13\nD. 192.168.1.14",
    "某工程师计划使用Zabbix监控企业中的服务器，在安装配置完Zabbix后，发现Zabbix Server无法启动，以下哪一项不是导致该问题的原因?\nA. 数据库密码配置错误\nB. 数据库未安装到Zabbix服务器中\nC. 数据库中未为 Zabbix创建库和用户\nD. 数据库中未导入Zabbix数据表",
    "某企业使用Nginx做了负载均衡，某工程师在日常运维过程中，发现Nginx从来没有向某个后端服务器转发过请求，以下哪一项是导致该问题的最可能原因?\nA. 该后端服务器被配置为down\nB. 该后端服务器被配置为weight=0\nC. 该后端服务器max_conns配置为1\nD. 该后端服务器没有配置weight",
    "以下关于互联网的描述，哪个选项是正确的?\nA. Nginx在互联网中可以作为web服务器端，成为万维网的一个节点\nB. Nginx在万维网中可以作为ftp服务器的反向代理，并与ftp服务器的数量一一对应\nC. ansible可以批量部署Nginx服务器，并对其进行实时监控\nD. 互联网上的资源需使用Nginx进行七层反向代理，才能被通过http方式访问到",
    "根据下图分析，以下哪个选项的描述是正确的?\n```\nHTTP/1.1 200 \nServer: CloudiMF \nDate: Thu, 01 Jun 2023 00:36:39 GMT \nContent-Type:text/plain;charset=UTF-8 \nTransfer-Encoding: chunked \nConnection: keep-alive \n1ubanops-gtrace-id: v+710197-1685579799386-199328934 \nJubanops-new-id:179450 \nAccess-Control-Allow-Origin:* \nX-Kong-Upstream-Latency: 18 \nX-Kong-Proxy-Latency: 1 \nExpires:Thu,08 Jun 2023 00:36:39 GMT \nCache-Control: max-age=604800 \nCache-Control: public \nX-Frame-Options: SAMEORIGIN \nX-Content-Type-Options: nosniff \nX-Download-Options:noopen \nX-XSS-Protection: 1; mode=block \nContent-Encoding: gzip\n```\nA. 图中展示的是HTTP请求报文头，此次请求连接正常\nB. 图中展示的是HTTP响应报文头，此次连接为长连接\nC. 图中展示的是HTTP请求报文头，此次请求的服务器使用的是HTTP版本为1.1\nD. 图中展示的是HTTP响应报文头，此次连接的服务器前端使用了keepalived进行负载均衡",
    "根据下图进行分析，哪个选项是正确的?\n```\nActive Internet connections (only servers)\nProto Recv-Q Send-Q Local Address           Foreign Address         State       PID/Program name    \ntcp        0      0 0.0.0.0:80              0.0.0.0:*               LISTEN      1234/nginx: master  \ntcp        0      0 0.0.0.0:22              0.0.0.0:*               LISTEN      567/sshd            \ntcp6       0      0 :::443                  :::*                    LISTEN      1234/nginx: master  \n```\nA. 当前服务器上运行了Nginx服务，可通过http对其进行访问\nB. 当前服务器上运行了Nginx服务,可通过https对其进行访问\nC. 当前服务器上运行了ssh服务，但禁止了root登录\nD. 当前服务器上运行了 ssh服务,仅能通过密钥方式登陆",
    "以下哪个场景适合使用 SAN作为两台服务器的共享存储?\nA. 两台服务器运行了基于 openEuler的 Apache集群\nB. 两台服务器运行了基于 openEuler的 OpenStack集群\nC. 两台服务器运行了基于 openEuler的 MySQL集群\nD. 两台服务器运行了基于 openEuler的 DNS集群",
    "以下关于存储及其使用方式的描述中，错误的是哪一项?\nA. 某医院计划新建病历影像库,适合使用GlusterFS作为存储\nB. 某银行建设交易系统数据库,适合使用 FC SAN作为存储\nC. 某公司互联网业务使用 Redis作为缓存,适合使用 IP SAN作为存储\nD. 某公司为员工配置虚拟机作为办公桌面，适合使用 FC SAN作为存储",
    "某工程师使用 Nginx的虚拟主机功能，对于同一台服务器上的不同资源，使用不同的端口号来区分,配置完成后发现不同端口访问的是同一个资源,那么可能存在的配置错误为以下哪一项?\nA. 关于虚拟主机的子配置文件 vhost.conf中 server模块的 root配置相同,且未配置 index\nB. 关于虚拟主机的子配置文件 vhost.conf中 server模块的 index配置相同\nC. 主配置文件conf.d中未包含虚拟主机相关的配置\nD. 主配置文件conf.d中也添加了相同的配置",
    "某工程师在配置 HAProxy时，制定了以下规则:\n```\nfrontend main\n    bind *:80\n    acl url_test path_end -i .jpg .gif .png .css .js\n    use_backend static if url_test\n    default_backend my_webserver\n```\n用户在访问www.test.com:8080/test.JPG时，会得到以下哪个响应?\nA. path_end对应的服务器返回的资源\nB. my_webserver对应的服务器返回的资源\nC. 资源不存在的报错信息\nD. 权限不足的报错信息",
    "某工程师需要使用Ansible命令 `ansible test -m yum -a \"name=httpd state=present\"` 在主机中成功安装了httpd服务，但用户去无法正常访问到默认页，以下哪个选项是可能原因?\nA. test主机中SELinux的状态为enforcing\nB. httpd服务未启动\nC. test主机的firewalld状态为disable\nD. 用户与test主机之间未配置免密登录",
    "以下哪个选项可以判断Nginx是否启用了负载均衡配置?\nA. `grep upstream /etc/nginx/nginx.conf /etc/nginx/conf.d/*`\nB. `grep location /etc/nginx/*`\nC. `grep server /etc/nginx/nginx.conf /etc/nginx/conf.d/*`\nD. `grep proxy_pass /etc/nginx/*`",
    "以下关于SAN存储的描述，哪个选项是正确的?\nA. SAN存储需依赖IP网络才能实现\nB. SAN存储属于集中式存储\nC. SAN存储属于分布式存储\nD. SAN存储仅工集群使用"
]

for i, question in enumerate(questions_single_choice, 1):
    doc.add_paragraph(f"{i}. {question}")

# 添加多选题部分
doc.add_heading('三、多选题', level=2)

questions_multiple_choice = [
    "(华为)某工程师使用Nginx为三台Apache服务器提供负载均衡功能，完成配置后，该工程师进行高可用性测试，发现每次访问都是同一台Apache服务器响应。以下哪些选项是造成这种现象的可能原因?\nA. 剩余两台服务器被添加了backup配置\nB. Nginx配置了IP Hash算法\nC. 剩余两台Apache服务器出现故障\nD. Nginx配置了虚拟主机",
    "(华为)某企业使用 Glusterfs分布式存储存放业务数据，该业务选择使用由4块Bricks组成的分散卷保存数据。在进行配置时，以下哪些选项不可作为disperse和redundancy的组合?\nA. disperse为 2, redundancy为 2\nB. disperse为 4, redundancy为 2\nC. disperse为4, redundancy为1\nD. disperse为 3, redundancy为 1",
    "某工程师使用LVS实现负载均衡功能，在选择工作模式时，该工程师不希望LVS服务器修改请求报文的大小，以下哪些工作模式可以满足此需求?\nA. TUN\nB. RR\nC. DR\nD. NAT",
    "(华为)某工程师使用 LVS(DR模式)实现 Nginx服务器的负载均衡功能，以下哪几项配置是正确的?\nA. DIP和RIP配置在不同网段\nB. LVS服务器和Nginx均配置VIP，且VIP一致\nC. RIP和VIP配置在同一网段\nD. LVS上VIP的网卡上关闭ARP广播和应答",
    "(华为)对下图中的信息分析，以下哪些描述是正确的?\n```\n[root@Gluster-01 ~]# pidstat | grep gluster \n10:57:34 AM     0     1156    0.06    0.01    0.00    0.00    0.07     0 glusterd\n10:57:34 AM     0     1402    0.00    0.01    0.00    0.00    0.01     2 glusterfsd\n10:57:34 AM     0     1431    0.00    0.00    0.00    0.00    0.01     0 glusterfsd\n10:57:34 AM     0     1461    0.11    0.01    0.00    0.00    0.12     0 glusterfs\n```\nA. 该主机至少有3个CPU\nB. glusterfs消耗了该主机的大量资源\nC. 该节点承载的glusterfs卷已被挂载\nD. 有两个gluster brick运行在该节点上",
    "(华为)某企业使用Nginx作为Web服务器,由于近期业务侧所有员工需要经常查看某个实时销售数据，访问量较大，在不增加服务器情况下，以下哪几项能够提高业务访问的效率?\nA. 在 nginx.conf中配置提高 types_hash_max_size的值\nB. 在nginx.conf中配置tcp_nodelay为off\nC. 在nginx.conf配置文件中提高 worker_process的值\nD. 在nginx.conf配置文件中提高 worker_connections的值",
    "用户在访问网站时收到了403的状态码，以下哪些选项是造成该现象的可能原因?\nA. web服务器由 Apache实现，且配置文件中将 Require not ip设置为该用户的IP地址\nB. web服务器由 Apache实现，且开启了SELinux和 firewalld安全措施\nC. web服务器由 Apache实现，并使用了 Prefork模式\nD. web服务器由 Apache实现，且配置文件中将 KeepAliveTimeout设置为0",
    "以下关于Apache和Nginx的区别，描述正确的是哪几项?\nA. Nginx相对于Apache配置更简洁，静态处理性能更高\nB. Apache默认使用同步多进程模型，一个连接对应一个进程，而Nginx是异步的，多个连接可以对应一个进程\nC. Nginx相对于Apache性能更加稳定\nD. Nginx相对于Apache更轻量，抗并发",
    "某企业基于Apache搭建网站服务，并通过LVS提供VIP对外提供服务。业务运行一段时间后，发现业务有时无法正常访问。以下哪些选项是造成这种现象的原因?\nA. 部分Apache服务器主配置文件中的主页存放路径错误，导致业务无法正常使用\nB. 环境未安装keepalived类的心跳检查功能，集群中有服务器故障\nC. Apache服务使用Prefork模式，某一时刻所有服务器子进程占用资源过多，服务器无法正常响应\nD. 集群LVS模式为NAT模式，未添加所有后端服务器，导致部分服务器无法访问",
    "(华为)某工程师使用LVS配置负载均衡时，采用DR的工作模式，更改配置前转发正常，更改配置后DR转发不正常。以下哪几项是该工程师需要在RS上检查的配置?\nA. 是否配置了VIP\nB. 是否在firewalld中放通了对应接口\nC. 是否打开了ip转发功能\nD. 是否关闭了arp广播",
    "(华为)以下哪几个服务建议配合Keepalived实现高可用?\nA. MySQL\nB. DNS\nC. Apache\nD. LVS",
    "根据下图中的输出分析，以下哪些选项是正确的?\n```\n[root@Nginx1 ~]# ps --forest --C nginx --o pid,ppid,cmd\n  PID  PPID CMD\n  970     1 nginx: master process /usr/sbin/nginx\n  973   970 \\_ nginx: worker process\n```\nA. 在命令中添加'%cpu'即可查看到对应进程的CPU占用率\nB. 该主机的物理核心数为1\nC. PID为973的父进程ID为970\nD. 该主机中的Nginx用来实现7层负载均衡",
    "在搭建LAMP集群时，至少需要安装和配置以下哪几种服务?\nA. GlusterFS\nB. MySQL\nC. Nginx\nD. Apache",
    "(华为)某工程师在Nginx的配置文件中添加了以下配置，关于该配置描述正确的是哪几个选项?\n```\nupstream image{\n    server 10.0.0.41:82;\n    server 10.0.0.42:82;\n}\n```\nA. 客户端访问两台server时，目的端口必须是82\nB. Nginx为两台server提供了负载均衡功能，使用的是默认的轮询加权算法\nC. 配置server模块，可实现对两台业务的四层或七层代理\nD. image是固定参数，不可随意更改",
    "(华为)以下关于存储架构的描述，正确的是哪几项?\nA. 元数据服务器限制了存储资源的横向扩展， GlusterFS具有高可扩展性，因此推测GlusterFS没有元数据服务器，属于去中心化架构。\nB. NAS具备容错性，支持多种方式保护数据完整性，同时具备易于访问的特性，因此推测NAS适用于企业存储和文件共享\nC. 小文件通常指大小在1MB以内的文件，GlusterFS要求预置Brick，因此推测GlusterFS比较适合存储小文件\nD. IP SAN基于TCP/IP网络进行数据传输， FC SAN需要专用光纤交换机进行数据传输，因此推测IP SAN存储的扩展性优于FC SAN存储",
    "(华为)某工程师希望使用 SaltStack维护公司内部GlusterFS集群，以下操作命令与其功能描述相符的是哪些项?\nA. 创建分布式复制卷 vol2: `salt gluster1 glusterfs.create vol2 '[\"gluster1:/export/vol2/brick\",\"gluster2:/export/vol2/brick\"]' replica=2 start=True`\nB. 在host1上创建新的brick，并挂载至newvolume: `salt 'glusterfs.add_volume_bricks newvolume host1:/brick`\nC. 启动 mycluster集群: `salt \"*\" glusterfs.start mycluster`\nD. 将host3添加进入存储信任池: `salt 'one.gluster.' glusterfs.peer host3`",
    "某大型企业IT环境包含各种服务器和网络设备，该企业要求实现自动化运维，在该场景中工程师可以使用Ansible管理以下哪些设备?\nA. openEuler服务器\nB. Windows Server 2012 R2服务器\nC. 防火墙\nD. 交换机",
    "某工程师在配置LVS时采用了LVS默认的工作模式，该模式支持以下哪些负载均衡算法?\nA. 轮询算法\nB. 源地址哈希算法\nC. 最少连接算法\nD. 最短期望延迟算法",
    "(华为)关于以下Keepalived配置文件描述正确的是哪几项?\n```\nvrrp_instance Nginx {\n    state BACKUP \n    interface ens3 \n    virtual_router_id 52 \n    priority 200 \n    advert_int 1 \n    authentication {\n        auth_type PASS \n        auth_pass 1111\n    }\n    virtual_ipaddress {\n        10.0.0.20/24\n    }\n}\n```\nA. 该配置是为Nginx服务器提供主备机制的\nB. 该配置为备服务器的配置，当主服务出现故障后，才可能升级为主服务\nC. 主备服务器的virtual_router_id需要保持不同,用以区分不同服务器\nD. 10.0.0.20是主备服务器的虚拟IP",
    "Nginx作为web服务器的反向代理时，会涉及到以下哪些选项的工作步骤?\nA. 客户端向Nginx发送请求\nB. Nginx在收到客户端发送的请求后，将请求转发到后端服务器\nC. 后端服务器直接将资源发送给客户端\nD. Nginx定期检查后端服务是否可用",
    "在openEuler中关闭SSH服务后，以下哪些服务或应用的使用会受影响?\nA. Ansible\nB. SaltStack\nC. HTTP\nD. Glusterfs",
    "(华为)将内核参数ip_forward设置为0以后，以下哪些服务的功能可能会受影响?\nA. Nginx\nB. LVS\nC. Keepalived\nD. iptables\nE. firewalld\nF. HAproxy",
    "对比Ansible和SaltStack的功能，以下描述哪些是正确的?\nA. Ansible的service_facts和SaltStack的pillar功能相似\nB. Ansible的yum和SaltStack的pkg功能相似\nC. Ansible的service和SaltStack的service功能相似\nD. Ansible的file和SaltStack的file功能相似",
    "根据下图信息进行分析，以下哪些选项是正确的?\n```\n/dev/sdb1 on /data type xfs (rw,relatime,seclabel,attr2,inode64,noquota)\n10.0.0.20:/wp on /data/wp type fuse.glusterfs (rw,relatime,user_id=0,group_id=0,default_permissions,allow_other,max_read=131072)\n```\nA. 10.0.0.20/wp是由Glusterfs提供的卷\nB. 10.0.0.20/wp会开机自动挂载到/data/wp目录\nC. 10.0.0.20/wp在被挂载时选用的是默认参数\nD. 10.0.0.20/wp可以通过nfs的方式挂载，但要求客户端支持并开启NFSv4",
    "(华为)某工程师在使用ansible playbook配置完成Keepalived后，发现集群中的所有主机都获取到了指定的VIP，查看日志后发现所有主机都首先将自己置为BACKUP，在未收到VIP消息通告后将自己升级为MASTER，并自动配置了VIP，以下哪些选项是造成这种现象的可能原因?\nA. VIP使用变量指定，但未生效\nB. keepalived的配置文件未使用jinja2模板，导致变量未生效\nC. 在playbook中未对firewalld进行配置\nD. 配置文件中未包含real_server相关配置\nE. 配置文件中未指定MASTER节点",
    "Apache和Nginx都支持虚拟主机功能，以下哪些选项是二者同时支持的区分不同主机的方式?\nA. IP地址\nB. 域名\nC. 端口号\nD. MAC地址",
    "在Zabbix中对一台提供反向代理服务的Nginx主机进行监控，以下哪些指标需要重点关注?\nA. 磁盘IO\nB. 磁盘容量\nC. CPU使用率\nD. 内存使用率",
    "某工程师使用Ansible在主机中安装Apache时，返回以下报错，哪些选项为造成该报错的可能原因?\n```\n10.0.0.218 | UNREACHABLE! => {\n    \"changed\": false,\n    \"unreachable\": true\n}\n```\nA. 10.0.0.218启用了firewalld且未放通ssh端口\nB. 10.0.0.218的ssh服务将 PasswordAuthentication设置为no\nC. 控制主机未配置对 10.0.0.218的免密登录\nD. 控制主机未开启 sshd服务",
    "(华为)某企业采用MySQL管理内部业务数据，工程师在选择数据库连接器时，可以选择以下哪些语言开发的程序?\nA. PHP\nB. Ruby\nC. JDBC\nD. Python",
    "(华为)某小型企业基于web服务器搭建了文件共享系统，员工可通过浏览器上传或下载部门文件，不同部门之间使用不同的路径来隔离，以下哪些配置可用于该系统的实现?\nA.\n```\n<VirtualHost *:81>\n    ServerName localhost\n    DocumentRoot \"/department1/\"\n    <Directory \"/department1/\">\n        AllowOverride None\n        Require all granted\n    </Directory>\n</VirtualHost>\n<VirtualHost *:82>\n    DocumentRoot \"/department2/\"\n    <Directory \"/department2/\">\n        AllowOverride None\n        Require all granted\n    </Directory>\n    ServerName localhost\n</VirtualHost>\n```\nB.\n```\nhosts: Nginx\nremote_user: root\ngather_facts: no\ntasks:\n    - name: create department1\n      file:\n          path: /department1\n          state: directory\n    - name: create department2\n      file:\n          path: /department2\n          state: directory\n```\nC.\n```\nfrontend main\n    bind *:80\n    acl url_department1 path_end department1\n    acl url_department2 path_end department2\n    use_backend department1 if url_department1\n    use_backend department2 if url_department2\n```\nD. （此处原文档无内容，保留空白）"
]

for i, question in enumerate(questions_multiple_choice, 1):
    doc.add_paragraph(f"{i}. {question}")

# 添加填空题部分
doc.add_heading('四、填空题', level=2)

questions_fill_blank = [
    "A-Ops包含多个组件，每个组件实现对应的功能，用户可根据自己的需求选择安装，其中需要部署在客户端组件是（ ）。",
    "（ ）起初是为LVS设计的，专门用来监控集群系统中各个服务节点的状态，两者搭配可以实现负载均衡和高可用。",
    "下图是（ ）代理的示意图。",
    "某LAMP架构的应用对浏览器类型有严格的要求,如果不满足就无法正常访问。管理员在该应用的前端使用（ ）来实现七层代理，将相符类型浏览器的请求转发至对应服务器，不符合要求的请求则响应为报错并提示安装正确浏览器的页面。",
    "使用Glusterfs创建一个分布式复制卷,该卷由12块Bricks组成,创建命令为:\n```\ngluster volume create gv3 node2:/ex/brick2 node1:/ex/brick1 node3:/ex/brick1 node1:/ex/brick2 node4:/ex/brick3 node2:/ex/brick1 node1:/ex/brick3 node3:/ex/brick2 node4:/ex/brick1 node4:/ex/brick2 node2:/ex/brick3 node3:/ex/brick3\n```\n客户端Client挂载该卷后，保存了多个文件在卷中，其中发现file1保存在node1:/ex/brick3中，那么根据Glusterfs的工作原理file1的另一副本保存在（ ）中。注意:考试时候卷的位置有点乱，要仔细分析。"
]

for i, question in enumerate(questions_fill_blank, 1):
    doc.add_paragraph(f"{i}. {question}")

# 保存文档
filename = "HCIP-openEulerV1.0-exam-客观题.docx"
doc.save(filename)
print(filename)